import io
import unittest
from datetime import datetime
from typing import TYPE_CHECKING
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from module_gat1400.dao.pass_record_dao import TruckPassRecordDao
from module_gat1400.entity.vo.pass_record_vo import PassRecordPageQueryModel
from module_gat1400.service.pass_record_service import PassRecordService
from module_truck.entity.do.truck_do import TruckVehicle

if TYPE_CHECKING:
    from sqlalchemy import ColumnElement


def compile_sql(statement) -> str:  # noqa: ANN001
    """将 SQLAlchemy 语句编译为 SQL 文本，便于断言条件。"""
    return str(statement.compile(compile_kwargs={'literal_binds': True}))


class PassRecordDaoTests(unittest.IsolatedAsyncioTestCase):
    async def test_get_list_applies_query_filters(self) -> None:
        db = AsyncMock()
        captured: dict = {}
        query = PassRecordPageQueryModel(
            plate_no='浙K',
            device_id='D001',
            start_time=datetime(2026, 8, 1),
            end_time=datetime(2026, 8, 17),
            vehicle_matched=True,
        )
        with patch(
            'module_gat1400.dao.pass_record_dao.PageUtil.paginate',
            AsyncMock(side_effect=lambda _db, stmt, _pn, _ps, _is_page: captured.update(stmt=stmt)),
        ):
            await TruckPassRecordDao.get_list(db, query, True, True)

        sql = compile_sql(captured['stmt'])
        self.assertIn('LIKE', sql)
        self.assertIn("'%浙K%'", sql)
        self.assertIn("'D001'", sql)
        self.assertIn('IS NOT NULL', sql)
        self.assertIn('pass_time', sql)

    async def test_get_list_scope_true_sees_all(self) -> None:
        """全量角色（scope=True）不追加数据权限条件。"""
        db = AsyncMock()
        captured: dict = {}
        with patch(
            'module_gat1400.dao.pass_record_dao.PageUtil.paginate',
            AsyncMock(side_effect=lambda _db, stmt, _pn, _ps, _is_page: captured.update(stmt=stmt)),
        ):
            await TruckPassRecordDao.get_list(db, PassRecordPageQueryModel(), True, True)

        sql = compile_sql(captured['stmt'])
        # 不包含任何 dept_id 条件
        self.assertNotIn('dept_id', sql)

    async def test_get_list_squad_scope_filters_by_vehicle_dept(self) -> None:
        """中队管理员（scope 作用于 truck_vehicle.dept_id）通过 outerjoin 过滤。"""
        db = AsyncMock()
        captured: dict = {}
        squad_dept_id = 9101
        scope: ColumnElement = TruckVehicle.dept_id == squad_dept_id
        with patch(
            'module_gat1400.dao.pass_record_dao.PageUtil.paginate',
            AsyncMock(side_effect=lambda _db, stmt, _pn, _ps, _is_page: captured.update(stmt=stmt)),
        ):
            await TruckPassRecordDao.get_list(db, PassRecordPageQueryModel(), scope, True)

        sql = compile_sql(captured['stmt'])
        self.assertIn('JOIN', sql.upper())
        self.assertIn('dept_id', sql)
        self.assertIn(str(squad_dept_id), sql)


class PassRecordExportTests(unittest.TestCase):
    def test_export_contains_mapped_columns(self) -> None:
        content = PassRecordService.export_pass_records(
            [
                {
                    'plateNo': '浙K12345',
                    'plateColor': '5',
                    'vehicleClass': '1',
                    'vehicleColor': '2',
                    'passTime': datetime(2026, 8, 17, 10, 30, 0),
                    'nameOfPassedRoad': '松阳大道',
                    'laneNo': 2,
                    'speed': 65,
                    'tollgateId': 'TG-1',
                    'deviceId': 'D001',
                    'vehicleMatched': True,
                    'deptName': '西屏中队',
                    'storageUrl1': 'http://x/1.jpg',
                }
            ]
        )
        worksheet = load_workbook(io.BytesIO(content), data_only=True).active

        self.assertEqual(worksheet.cell(2, 1).value, '浙K12345')
        # GA/T1400 色号 → 中文映射（5=蓝色）
        self.assertEqual(worksheet.cell(2, 2).value, '蓝色')
        self.assertEqual(worksheet.cell(2, 5).value, '2026-08-17 10:30:00')
        self.assertEqual(worksheet.cell(2, 8).value, 65)
        self.assertEqual(worksheet.cell(2, 12).value, '是')
        self.assertEqual(worksheet.cell(2, 13).value, '西屏中队')


if __name__ == '__main__':
    unittest.main()
