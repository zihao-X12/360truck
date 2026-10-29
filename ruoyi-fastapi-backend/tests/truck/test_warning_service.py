import io
import unittest

from openpyxl import load_workbook

from module_truck.service.warning_service import TruckWarningService


class TruckWarningServiceTests(unittest.TestCase):
    def test_default_rule_only_warns_when_nothing_installed(self) -> None:
        """默认阈值 2/2：应安装 2 项，只有一项都没装才预警。"""
        required, level_one, level_two = TruckWarningService.parse_rule(None)

        self.assertEqual(required, ['brake_assist', 'panoramic_360'])
        self.assertEqual((level_one, level_two), (2, 2))

    def test_default_thresholds_warn_only_when_nothing_installed(self) -> None:
        """缺装 2 项 → 一级预警；缺装 1 项与 0 项均不预警。"""
        self.assertEqual(TruckWarningService.calculate_level(2, 2, 2), 'level_one')
        self.assertEqual(TruckWarningService.calculate_level(1, 2, 2), 'none')
        self.assertEqual(TruckWarningService.calculate_level(0, 2, 2), 'none')

    def test_two_level_grading_still_supported_by_thresholds(self) -> None:
        """阈值机制保留两级能力：给出 3/1 仍能区分一级与二级，便于后续恢复分级。"""
        self.assertEqual(TruckWarningService.calculate_level(3, 3, 1), 'level_one')
        self.assertEqual(TruckWarningService.calculate_level(2, 3, 1), 'level_two')
        self.assertEqual(TruckWarningService.calculate_level(1, 3, 1), 'level_two')
        self.assertEqual(TruckWarningService.calculate_level(0, 3, 1), 'none')

    def test_warning_export_contains_missing_device_names(self) -> None:
        content = TruckWarningService.export_warnings(
            [
                {
                    'plateNo': '浙K12345',
                    'warningLevel': 'level_one',
                    'missingDeviceTypes': ['brake_assist', 'panoramic_360'],
                }
            ]
        )
        worksheet = load_workbook(io.BytesIO(content), data_only=True).active

        self.assertEqual(worksheet.cell(2, 1).value, '浙K12345')
        self.assertEqual(worksheet.cell(2, 7).value, '一级预警')
        self.assertIn('刹车辅助系统', worksheet.cell(2, 8).value)


if __name__ == '__main__':
    unittest.main()
