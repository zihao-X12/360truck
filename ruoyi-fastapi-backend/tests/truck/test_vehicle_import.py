import io
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import Workbook, load_workbook

from module_truck.service.vehicle_import_service import TruckVehicleImportService


class TruckVehicleImportTests(unittest.TestCase):
    @staticmethod
    def _workbook_bytes(rows: list[list]) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        for row in rows:
            worksheet.append(row)
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_legacy_headers_distinguish_two_phone_columns(self) -> None:
        content = self._workbook_bytes([
            TruckVehicleImportService.LEGACY_HEADERS,
            ['工地A', ' 浙a 12345 ', '自卸货车', None, None, '张三', '13800000000', '李四', '13900000000', '无', None, '是', '否', '西屏中队', '是'],
        ])

        headers, rows = TruckVehicleImportService.parse_workbook(content)

        self.assertEqual(headers, TruckVehicleImportService.LEGACY_HEADERS)
        self.assertEqual(rows[0]['values']['owner_phone'], '13800000000')
        self.assertEqual(rows[0]['values']['driver_phone'], '13900000000')

    def test_build_vehicle_maps_defaults_and_devices(self) -> None:
        values = {
            'plate_no': ' 浙a 12345 ',
            'vehicle_type': '自卸货车',
            'dept_name': '西屏中队',
            'source_type': '是',
            'brake_assist': '是',
            'panoramic_360': '否',
        }

        vehicle = TruckVehicleImportService.build_vehicle(
            values,
            {'西屏中队': 9101},
            {'自卸货车': 'dump_truck'},
            {'黄牌': 'yellow', 'yellow': 'yellow'},
            'admin',
        )

        self.assertEqual(vehicle.plate_no, '浙A12345')
        self.assertEqual(vehicle.plate_color, 'yellow')
        self.assertEqual(vehicle.source_type, 'local')
        self.assertEqual([device.device_type for device in vehicle.devices], ['brake_assist'])

    def test_invalid_source_type_has_real_value_in_message(self) -> None:
        with self.assertRaisesRegex(ValueError, '未知'):
            TruckVehicleImportService.build_vehicle(
                {'plate_no': '浙A12345', 'dept_name': '西屏中队', 'source_type': '未知'},
                {'西屏中队': 9101},
                {},
                {'yellow': 'yellow'},
                'admin',
            )

    def test_error_workbook_contains_row_and_message(self) -> None:
        row = {'sheet': 'Sheet1', 'row_number': 2, 'raw': ['工地A', None]}
        content = TruckVehicleImportService.build_error_workbook(['工地名称', '车牌'], [(row, '车牌不能为空')])
        worksheet = load_workbook(io.BytesIO(content), data_only=True).active

        self.assertEqual(worksheet.cell(2, 1).value, 'Sheet1')
        self.assertEqual(worksheet.cell(2, 2).value, 2)
        self.assertEqual(worksheet.cell(2, 5).value, '车牌不能为空')

    def test_empty_export_keeps_business_headers(self) -> None:
        content = TruckVehicleImportService.export_vehicles([])
        worksheet = load_workbook(io.BytesIO(content), data_only=True).active
        headers = [cell.value for cell in worksheet[1]]

        self.assertIn('辖区中队', headers)
        self.assertIn('是否安装刹辅系统', headers)


class TruckVehicleOwnerImportTests(unittest.IsolatedAsyncioTestCase):
    @staticmethod
    def _row(plate_no: str, **values) -> dict:
        return {
            'sheet': 'Sheet1',
            'row_number': 2,
            'raw': [],
            'values': {'plate_no': plate_no, 'dept_name': '西屏中队', **values},
        }

    async def _import_row(self, row: dict, owner_cache: dict | None = None) -> tuple[str | None, AsyncMock, AsyncMock]:
        owner_cache = owner_cache if owner_cache is not None else {}
        with (
            patch('module_truck.service.vehicle_import_service.TruckOwnerDao.get_by_name', new_callable=AsyncMock) as get_owner,
            patch('module_truck.service.vehicle_import_service.TruckOwnerDao.add', new_callable=AsyncMock) as add_owner,
            patch('module_truck.service.vehicle_import_service.TruckVehicleService.add', new_callable=AsyncMock) as add_vehicle,
        ):
            get_owner.return_value = None
            add_owner.return_value = SimpleNamespace(owner_id=1001)
            error = await TruckVehicleImportService._import_row(
                AsyncMock(), row, {'西屏中队': 9101}, {}, {'yellow': 'yellow'}, 'admin', set(), set(), owner_cache
            )
            return error, add_owner, add_vehicle

    async def test_motor_owner_creates_individual_and_links_vehicle(self) -> None:
        error, add_owner, add_vehicle = await self._import_row(
            self._row('浙A12345', motor_owner_name='  张三  ', owner_phone='13800000000')
        )

        self.assertIsNone(error)
        owner = add_owner.await_args.args[1]
        vehicle = add_vehicle.await_args.args[1]
        self.assertEqual((owner.owner_type, owner.owner_name, owner.contact_phone), ('individual', '张三', '13800000000'))
        self.assertEqual((vehicle.owner_id, vehicle.owner_type), (1001, 'individual'))

    async def test_affiliated_company_is_used_when_motor_owner_is_empty(self) -> None:
        error, add_owner, add_vehicle = await self._import_row(
            self._row('浙A12346', affiliated_company='松阳运输公司')
        )

        self.assertIsNone(error)
        owner = add_owner.await_args.args[1]
        vehicle = add_vehicle.await_args.args[1]
        self.assertEqual((owner.owner_type, owner.owner_name), ('enterprise', '松阳运输公司'))
        self.assertEqual((vehicle.owner_id, vehicle.owner_type), (1001, 'enterprise'))

    async def test_empty_owner_fields_do_not_create_or_link_owner(self) -> None:
        error, add_owner, add_vehicle = await self._import_row(self._row('浙A12347'))

        self.assertIsNone(error)
        add_owner.assert_not_awaited()
        vehicle = add_vehicle.await_args.args[1]
        self.assertIsNone(vehicle.owner_id)
        self.assertIsNone(vehicle.owner_type)

    async def test_successful_owner_is_reused_from_batch_cache(self) -> None:
        owner_cache: dict[tuple[str, str], int] = {('individual', '张三'): 1001}
        error, add_owner, add_vehicle = await self._import_row(
            self._row('浙A12348', motor_owner_name='张三'), owner_cache
        )

        self.assertIsNone(error)
        add_owner.assert_not_awaited()
        vehicle = add_vehicle.await_args.args[1]
        self.assertEqual(vehicle.owner_id, 1001)

    async def test_existing_owner_is_reused_without_creating_duplicate(self) -> None:
        with (
            patch('module_truck.service.vehicle_import_service.TruckOwnerDao.get_by_name', new_callable=AsyncMock) as get_owner,
            patch('module_truck.service.vehicle_import_service.TruckOwnerDao.add', new_callable=AsyncMock) as add_owner,
            patch('module_truck.service.vehicle_import_service.TruckVehicleService.add', new_callable=AsyncMock) as add_vehicle,
        ):
            get_owner.return_value = SimpleNamespace(owner_id=2001)
            error = await TruckVehicleImportService._import_row(
                AsyncMock(), self._row('浙A12349', motor_owner_name='张三'),
                {'西屏中队': 9101}, {}, {'yellow': 'yellow'}, 'admin', set(), set(), {}
            )

        self.assertIsNone(error)
        get_owner.assert_awaited_once_with(unittest.mock.ANY, 'individual', '张三')
        add_owner.assert_not_awaited()
        self.assertEqual(add_vehicle.await_args.args[1].owner_id, 2001)

    async def test_failed_vehicle_does_not_cache_rolled_back_owner(self) -> None:
        owner_cache: dict[tuple[str, str], int] = {}
        with (
            patch('module_truck.service.vehicle_import_service.TruckOwnerDao.get_by_name', new_callable=AsyncMock) as get_owner,
            patch('module_truck.service.vehicle_import_service.TruckOwnerDao.add', new_callable=AsyncMock) as add_owner,
            patch('module_truck.service.vehicle_import_service.TruckVehicleService.add', new_callable=AsyncMock) as add_vehicle,
        ):
            get_owner.return_value = None
            add_owner.return_value = SimpleNamespace(owner_id=3001)
            add_vehicle.side_effect = ValueError('车辆写入失败')
            error = await TruckVehicleImportService._import_row(
                AsyncMock(), self._row('浙A12350', motor_owner_name='张三'),
                {'西屏中队': 9101}, {}, {'yellow': 'yellow'}, 'admin', set(), set(), owner_cache
            )

        self.assertEqual(error, '车辆写入失败')
        self.assertEqual(owner_cache, {})


if __name__ == '__main__':
    unittest.main()
