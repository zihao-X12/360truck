import base64
import io
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any
from uuid import uuid4

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError
from redis.asyncio import Redis
from sqlalchemy import ColumnElement, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

from exceptions.exception import ServiceException
from module_admin.entity.do.dept_do import SysDept
from module_admin.entity.do.dict_do import SysDictData
from module_truck.dao.owner_dao import TruckOwnerDao
from module_truck.entity.vo.owner_vo import TruckOwnerModel
from module_truck.entity.vo.vehicle_vo import TruckVehicleDeviceModel, TruckVehicleModel
from module_truck.service.vehicle_service import TruckVehicleService
from utils.excel_util import ExcelUtil


class TruckVehicleImportService:
    MAX_ROWS = 500
    ERROR_CACHE_SECONDS = 1800
    ERROR_CACHE_PREFIX = 'truck:vehicle:import:error'
    PHONE_PATTERN = re.compile(r'^1\d{10}$')
    LEGACY_HEADERS = [
        '工地名称', '车牌', '车辆类型', '车辆品牌', '保险情况', '机动车所有人', '联系电话',
        '驾驶员', '联系电话', '挂靠公司', '保险公司及金额', '是否安装刹辅系统', '360',
        '辖区中队', '是否本地车辆',
    ]
    TEMPLATE_HEADERS = [
        '工地名称', '车牌', '车牌颜色', '车辆类型', '车辆品牌', '车架号', '核定载重', '保险情况',
        '机动车所有人', '车主联系电话', '驾驶员', '驾驶员联系电话', '挂靠公司', '保险公司及金额',
        '是否安装刹辅系统', '360全景影像', '辖区中队', '是否本地车辆',
    ]

    @staticmethod
    def _text(value: Any) -> str | None:
        if value is None:
            return None
        result = str(value).strip()
        return result if result else None

    @classmethod
    def _normalize_owner_name(cls, value: Any) -> str | None:
        text = cls._text(value)
        return ' '.join(text.split()) if text else None

    @classmethod
    def _owner_identity(cls, vehicle: TruckVehicleModel) -> tuple[str, str] | None:
        owner_name = cls._normalize_owner_name(vehicle.motor_owner_name)
        if owner_name:
            return 'individual', owner_name
        company_name = cls._normalize_owner_name(vehicle.affiliated_company)
        if company_name:
            return 'enterprise', company_name
        return None

    @classmethod
    async def _resolve_or_create_owner(
        cls,
        db: AsyncSession,
        vehicle: TruckVehicleModel,
        owner_phone: str | None,
        user_name: str,
        owner_cache: dict[tuple[str, str], int],
    ) -> tuple[tuple[str, str], int] | None:
        identity = cls._owner_identity(vehicle)
        if not identity:
            return None
        owner_type, owner_name = identity
        cache_key = (owner_type, owner_name.casefold())
        owner_id = owner_cache.get(cache_key)
        if owner_id is None:
            owner = await TruckOwnerDao.get_by_name(db, owner_type, owner_name)
            if owner is None:
                now = datetime.now()
                owner = await TruckOwnerDao.add(
                    db,
                    TruckOwnerModel(
                        ownerType=owner_type,
                        ownerName=owner_name,
                        deptId=vehicle.dept_id,
                        contactPhone=owner_phone if owner_type == 'individual' else None,
                        createBy=user_name,
                        updateBy=user_name,
                        createTime=now,
                        updateTime=now,
                    ),
                )
            owner_id = owner.owner_id
        vehicle.owner_id = owner_id
        vehicle.owner_type = owner_type
        return cache_key, owner_id

    @classmethod
    def _header_indexes(cls, headers: list[str | None]) -> dict[str, int]:
        indexes: dict[str, int] = {}
        phone_count = 0
        aliases = {
            '工地名称': 'project_name', '车牌': 'plate_no', '车牌号': 'plate_no', '车牌颜色': 'plate_color',
            '车辆类型': 'vehicle_type', '车辆品牌': 'vehicle_brand', '车架号': 'vin', 'VIN': 'vin',
            '核定载重': 'rated_load', '保险情况': 'insurance_status', '机动车所有人': 'motor_owner_name',
            '车主联系电话': 'owner_phone', '驾驶员': 'driver_name', '驾驶员联系电话': 'driver_phone',
            '挂靠公司': 'affiliated_company', '保险公司及金额': 'insurance_info',
            '是否安装刹辅系统': 'brake_assist', '360': 'panoramic_360', '360全景影像': 'panoramic_360',
            '辖区中队': 'dept_name', '是否本地车辆': 'source_type',
        }
        for index, raw_header in enumerate(headers):
            header = cls._text(raw_header)
            if header == '联系电话':
                key = 'owner_phone' if phone_count == 0 else 'driver_phone'
                phone_count += 1
            else:
                key = aliases.get(header or '')
            if key and key not in indexes:
                indexes[key] = index
        return indexes

    @classmethod
    def parse_workbook(cls, content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ServiceException(message='无法读取Excel文件，请确认文件未损坏且格式为xlsx') from exc

        parsed_rows: list[dict[str, Any]] = []
        output_headers: list[str] = []
        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows or not any(cls._text(value) for row in rows for value in row):
                continue
            first_row = [cls._text(value) for value in rows[0]]
            has_header = '车牌' in first_row or '车牌号' in first_row
            headers = first_row if has_header else cls.LEGACY_HEADERS
            indexes = cls._header_indexes(headers)
            if 'plate_no' not in indexes:
                raise ServiceException(message=f'工作表“{worksheet.title}”缺少“车牌”列')
            if not output_headers:
                output_headers = [header or f'第{index + 1}列' for index, header in enumerate(headers)]
            start_row = 1 if has_header else 0
            for row_number, row in enumerate(rows[start_row:], start=start_row + 1):
                values = list(row)
                if not any(cls._text(value) for value in values):
                    continue
                parsed_rows.append({
                    'sheet': worksheet.title,
                    'row_number': row_number,
                    'raw': values,
                    'values': {key: values[index] if index < len(values) else None for key, index in indexes.items()},
                })
        if not parsed_rows:
            raise ServiceException(message='Excel中没有可导入的数据行')
        if len(parsed_rows) > cls.MAX_ROWS:
            raise ServiceException(message=f'单次最多导入{cls.MAX_ROWS}行，当前共{len(parsed_rows)}行')
        return output_headers, parsed_rows

    @classmethod
    def _map_option(cls, value: Any, options: dict[str, str], field_name: str, default: str | None = None) -> str | None:
        text = cls._text(value)
        if not text:
            return default
        mapped = options.get(text.lower()) or options.get(text)
        if not mapped:
            raise ValueError(f'{field_name}“{text}”未配置')
        return mapped

    @classmethod
    def _map_source_type(cls, value: Any) -> str | None:
        text = cls._text(value)
        if not text:
            return None
        mapping = {'是': 'local', '本地': 'local', '本地车辆': 'local', 'local': 'local', '否': 'nonlocal',
                   '外地': 'nonlocal', '外地车辆': 'nonlocal', 'nonlocal': 'nonlocal'}
        if text.lower() not in mapping and text not in mapping:
            raise ValueError(f'是否本地车辆“{text}”无法识别，应填写是/否或本地/外地')
        return mapping.get(text.lower()) or mapping[text]

    @classmethod
    def _device_enabled(cls, value: Any, field_name: str) -> bool:
        text = cls._text(value)
        if not text:
            return False
        if text.lower() in {'是', '有', '已安装', '1', 'yes', 'y'}:
            return True
        if text.lower() in {'否', '无', '未安装', '0', 'no', 'n'}:
            return False
        raise ValueError(f'{field_name}“{text}”无法识别，应填写是或否')

    @classmethod
    def build_vehicle(
        cls,
        values: dict[str, Any],
        departments: dict[str, int],
        vehicle_types: dict[str, str],
        plate_colors: dict[str, str],
        user_name: str,
    ) -> TruckVehicleModel:
        plate_no = cls._text(values.get('plate_no'))
        if not plate_no:
            raise ValueError('车牌不能为空')
        dept_name = cls._text(values.get('dept_name'))
        if not dept_name:
            raise ValueError('辖区中队不能为空')
        dept_id = departments.get(dept_name)
        if not dept_id:
            raise ValueError(f'辖区中队“{dept_name}”不存在或无权使用')
        owner_phone = cls._text(values.get('owner_phone'))
        driver_phone = cls._text(values.get('driver_phone'))
        for label, phone in [('车主联系电话', owner_phone), ('驾驶员联系电话', driver_phone)]:
            if phone and not cls.PHONE_PATTERN.fullmatch(phone):
                raise ValueError(f'{label}“{phone}”格式错误，应为11位手机号')
        rated_load = cls._text(values.get('rated_load'))
        try:
            rated_load_value = Decimal(rated_load) if rated_load else None
        except InvalidOperation as exc:
            raise ValueError(f'核定载重“{rated_load}”不是有效数字') from exc
        devices = []
        if cls._device_enabled(values.get('brake_assist'), '是否安装刹辅系统'):
            devices.append(TruckVehicleDeviceModel(deviceType='brake_assist'))
        if cls._device_enabled(values.get('panoramic_360'), '360全景影像'):
            devices.append(TruckVehicleDeviceModel(deviceType='panoramic_360'))
        now = datetime.now()
        vehicle = TruckVehicleModel(
            projectName=cls._text(values.get('project_name')),
            plateNo=plate_no,
            plateColor=cls._map_option(values.get('plate_color'), plate_colors, '车牌颜色', 'yellow'),
            vehicleType=cls._map_option(values.get('vehicle_type'), vehicle_types, '车辆类型'),
            vehicleBrand=cls._text(values.get('vehicle_brand')),
            vin=cls._text(values.get('vin')),
            ratedLoad=rated_load_value,
            sourceType=cls._map_source_type(values.get('source_type')),
            motorOwnerName=cls._text(values.get('motor_owner_name')),
            driverName=cls._text(values.get('driver_name')),
            driverPhone=driver_phone,
            affiliatedCompany=cls._text(values.get('affiliated_company')),
            insuranceStatus=cls._text(values.get('insurance_status')),
            insuranceInfo=cls._text(values.get('insurance_info')),
            deptId=dept_id,
            devices=devices,
            createBy=user_name,
            updateBy=user_name,
            createTime=now,
            updateTime=now,
        )
        vehicle.validate_fields()
        return vehicle

    @staticmethod
    def _validation_message(exc: Exception) -> str:
        if isinstance(exc, ServiceException):
            return exc.message or '业务校验失败'
        if isinstance(exc, ValidationError):
            return '；'.join(str(item['msg']).replace('Value error, ', '') for item in exc.errors())
        if isinstance(exc, SQLAlchemyError):
            return '数据库处理失败，请联系管理员检查服务日志'
        return str(exc) or '导入失败'

    @classmethod
    async def _options(cls, db: AsyncSession, dict_type: str) -> dict[str, str]:
        records = (await db.execute(select(SysDictData).where(SysDictData.dict_type == dict_type, SysDictData.status == '0'))).scalars().all()
        result: dict[str, str] = {}
        for record in records:
            result[record.dict_label] = record.dict_value
            result[record.dict_value.lower()] = record.dict_value
        return result

    @classmethod
    async def _import_row(
        cls,
        db: AsyncSession,
        row: dict[str, Any],
        departments: dict[str, int],
        vehicle_types: dict[str, str],
        plate_colors: dict[str, str],
        user_name: str,
        seen_plates: set[tuple[str, str]],
        seen_vins: set[str],
        owner_cache: dict[tuple[str, str], int],
    ) -> str | None:
        try:
            vehicle = cls.build_vehicle(row['values'], departments, vehicle_types, plate_colors, user_name)
            plate_key = (vehicle.plate_no or '', vehicle.plate_color or '')
            if plate_key in seen_plates:
                raise ValueError(f'文件内车牌号{vehicle.plate_no}和车牌颜色重复')
            if vehicle.vin and vehicle.vin in seen_vins:
                raise ValueError(f'文件内车架号{vehicle.vin}重复')
            resolved_owner = await cls._resolve_or_create_owner(
                db,
                vehicle,
                cls._text(row['values'].get('owner_phone')),
                user_name,
                owner_cache,
            )
            await TruckVehicleService.add(db, vehicle)
            if resolved_owner:
                cache_key, owner_id = resolved_owner
                owner_cache[cache_key] = owner_id
            seen_plates.add(plate_key)
            if vehicle.vin:
                seen_vins.add(vehicle.vin)
            return None
        except Exception as exc:
            return cls._validation_message(exc)

    @classmethod
    async def import_file(
        cls,
        db: AsyncSession,
        redis: Redis,
        file: UploadFile,
        user_name: str,
        dept_scope: ColumnElement,
    ) -> dict[str, Any]:
        filename = file.filename or ''
        if not filename.lower().endswith('.xlsx'):
            raise ServiceException(message='仅支持xlsx格式的Excel文件')
        content = await file.read()
        headers, rows = cls.parse_workbook(content)
        departments = {
            record.dept_name: record.dept_id
            for record in (await db.execute(select(SysDept).where(SysDept.del_flag == '0', SysDept.status == '0', dept_scope))).scalars().all()
        }
        vehicle_types = await cls._options(db, 'truck_vehicle_type')
        plate_colors = await cls._options(db, 'truck_plate_color')
        errors: list[tuple[dict[str, Any], str]] = []
        seen_plates: set[tuple[str, str]] = set()
        seen_vins: set[str] = set()
        owner_cache: dict[tuple[str, str], int] = {}
        success_count = 0
        for row in rows:
            error = await cls._import_row(
                db, row, departments, vehicle_types, plate_colors, user_name, seen_plates, seen_vins, owner_cache
            )
            if error:
                errors.append((row, error))
            else:
                success_count += 1
        batch_id = None
        if errors:
            batch_id = uuid4().hex
            error_bytes = cls.build_error_workbook(headers, errors)
            await redis.setex(f'{cls.ERROR_CACHE_PREFIX}:{batch_id}', cls.ERROR_CACHE_SECONDS, base64.b64encode(error_bytes))
        return {'successCount': success_count, 'errorCount': len(errors), 'batchId': batch_id}

    @classmethod
    def build_error_workbook(cls, headers: list[str], errors: list[tuple[dict[str, Any], str]]) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '导入错误'
        output_headers = ['工作表', '原始行号', *headers, '错误信息']
        worksheet.append(output_headers)
        for row, message in errors:
            raw = list(row['raw'])[:len(headers)]
            raw.extend([None] * (len(headers) - len(raw)))
            worksheet.append([row['sheet'], row['row_number'], *raw, message])
        header_fill = PatternFill('solid', fgColor='B42318')
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        worksheet.freeze_panes = 'A2'
        worksheet.auto_filter.ref = worksheet.dimensions
        for index, header in enumerate(output_headers, 1):
            worksheet.column_dimensions[get_column_letter(index)].width = 18 if header != '错误信息' else 42
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    @classmethod
    def get_template(cls) -> bytes:
        selector_headers = ['车牌颜色', '车辆类型', '是否安装刹辅系统', '360全景影像', '辖区中队', '是否本地车辆']
        options = [
            {'车牌颜色': ['黄牌', '蓝牌', '绿牌', '其他']},
            {'车辆类型': ['自卸货车', '混凝土搅拌车', '牵引车', '载货汽车', '其他']},
            {'是否安装刹辅系统': ['是', '否']}, {'360全景影像': ['是', '否']},
            {'辖区中队': ['西屏中队', '古市中队', '江南中队', '象溪中队', '大东坝中队', '玉岩中队']},
            {'是否本地车辆': ['是', '否']},
        ]
        return ExcelUtil.get_excel_template(cls.TEMPLATE_HEADERS, selector_headers, options)

    @classmethod
    def export_vehicles(cls, vehicles: list[dict[str, Any]]) -> bytes:
        mapping = {
            'projectName': '工地名称', 'plateNo': '车牌', 'plateColor': '车牌颜色', 'vehicleType': '车辆类型',
            'vehicleBrand': '车辆品牌', 'vin': '车架号', 'ratedLoad': '核定载重', 'insuranceStatus': '保险情况',
            'motorOwnerName': '机动车所有人', 'driverName': '驾驶员', 'driverPhone': '驾驶员联系电话',
            'affiliatedCompany': '挂靠公司', 'insuranceInfo': '保险公司及金额',
            'brakeAssist': '是否安装刹辅系统', 'panoramic360': '360全景影像', 'deptName': '辖区中队',
            'sourceType': '来源', 'vehicleStatus': '车辆状态', 'supervisionFlag': '重点监管',
            'warningLevel': '预警等级', 'missingDeviceCount': '缺装设备数', 'createTime': '登记时间',
        }
        labels = {
            'yellow': '黄牌', 'blue': '蓝牌', 'green': '绿牌', 'other': '其他',
            'dump_truck': '自卸货车', 'mixer_truck': '混凝土搅拌车', 'tractor': '牵引车', 'cargo_truck': '载货汽车',
            'local': '本地车辆', 'nonlocal': '外地车辆', 'normal': '正常', 'stopped': '停运',
            'transferred': '转出', 'scrapped': '报废', 'level_one': '一级预警', 'level_two': '二级预警', 'none': '无预警',
        }
        rows = []
        for vehicle in vehicles:
            item = dict(vehicle)
            for field in ('plateColor', 'vehicleType', 'sourceType', 'vehicleStatus', 'warningLevel'):
                item[field] = labels.get(item.get(field), item.get(field))
            item['supervisionFlag'] = '是' if item.get('supervisionFlag') == '1' else '否'
            rows.append(item)
        return ExcelUtil.export_list2excel(rows, mapping)
