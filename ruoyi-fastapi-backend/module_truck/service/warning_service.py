import io
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import ColumnElement, select
from sqlalchemy.ext.asyncio import AsyncSession

from common.vo import CrudResponseModel
from module_admin.entity.do.dept_do import SysDept
from module_truck.dao.vehicle_dao import TruckVehicleDao
from module_truck.dao.warning_dao import TruckWarningDao
from module_truck.entity.do.truck_do import TruckOwner, TruckWarningRule
from module_truck.entity.vo.warning_vo import TruckWarningPageQueryModel


class TruckWarningService:
    DEFAULT_REQUIRED_DEVICE_TYPES = ['brake_assist', 'panoramic_360']
    DEVICE_NAMES = {'brake_assist': '刹车辅助系统', 'panoramic_360': '360全景影像'}
    LEVEL_NAMES = {'level_one': '一级预警', 'level_two': '二级预警', 'none': '无预警'}

    @classmethod
    def parse_rule(cls, rule: TruckWarningRule | None) -> tuple[list[str], int, int]:
        if not rule:
            return cls.DEFAULT_REQUIRED_DEVICE_TYPES, 2, 1
        try:
            required = json.loads(rule.required_device_types or '[]')
        except (TypeError, json.JSONDecodeError):
            required = []
        required = [item for item in required if isinstance(item, str)]
        return required or cls.DEFAULT_REQUIRED_DEVICE_TYPES, rule.level_one_threshold, rule.level_two_threshold

    @staticmethod
    def calculate_level(missing_count: int, level_one_threshold: int, level_two_threshold: int) -> str:
        if missing_count >= level_one_threshold:
            return 'level_one'
        if missing_count >= level_two_threshold:
            return 'level_two'
        return 'none'

    @classmethod
    async def recalculate(cls, db: AsyncSession, scope: ColumnElement) -> CrudResponseModel:
        rule = await TruckWarningDao.get_enabled_rule(db)
        required, level_one_threshold, level_two_threshold = cls.parse_rule(rule)
        vehicles = await TruckWarningDao.get_active_vehicles(db, scope)
        devices = await TruckVehicleDao.get_device_types(db, [vehicle.vehicle_id for vehicle in vehicles])
        try:
            for vehicle in vehicles:
                missing_count = len(set(required) - devices.get(vehicle.vehicle_id, set()))
                await TruckWarningDao.update_warning(
                    db,
                    vehicle.vehicle_id,
                    missing_count,
                    cls.calculate_level(missing_count, level_one_threshold, level_two_threshold),
                )
            await db.commit()
        except Exception:
            await db.rollback()
            raise
        return CrudResponseModel(is_success=True, message=f'预警重算完成，共处理{len(vehicles)}辆车', result={'count': len(vehicles)})

    @classmethod
    async def get_list(
        cls, db: AsyncSession, query: TruckWarningPageQueryModel, scope: ColumnElement, is_page: bool
    ) -> Any:
        result = await TruckWarningDao.get_list(db, query, scope, is_page)
        rows = result.rows if is_page else result
        vehicle_ids = [row['vehicleId'] for row in rows]
        device_types = await TruckVehicleDao.get_device_types(db, vehicle_ids)
        rule = await TruckWarningDao.get_enabled_rule(db)
        required, _, _ = cls.parse_rule(rule)
        dept_names = dict((await db.execute(select(SysDept.dept_id, SysDept.dept_name))).all())
        owner_ids = [row['ownerId'] for row in rows if row.get('ownerId')]
        owner_names = {}
        if owner_ids:
            owner_names = dict(
                (await db.execute(select(TruckOwner.owner_id, TruckOwner.owner_name).where(TruckOwner.owner_id.in_(owner_ids)))).all()
            )
        enriched = []
        for row in rows:
            installed = device_types.get(row['vehicleId'], set())
            enriched.append(
                {
                    'vehicleId': row['vehicleId'],
                    'plateNo': row['plateNo'],
                    'plateColor': row['plateColor'],
                    'vehicleType': row.get('vehicleType'),
                    'projectName': row.get('projectName'),
                    'deptId': row['deptId'],
                    'deptName': dept_names.get(row['deptId']),
                    'ownerName': owner_names.get(row.get('ownerId')),
                    'warningLevel': row['warningLevel'],
                    'missingDeviceCount': row['missingDeviceCount'],
                    'missingDeviceTypes': [item for item in required if item not in installed],
                }
            )
        if is_page:
            result.rows = enriched
            return result
        return enriched

    @classmethod
    def export_warnings(cls, rows: list[dict]) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '重点监管预警'
        headers = ['车牌号', '车牌颜色', '车辆类型', '工地名称', '所属中队', '归属主体', '预警等级', '缺失设备']
        worksheet.append(headers)
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='2F6B4F')
        for row in rows:
            worksheet.append(
                [
                    row.get('plateNo'),
                    row.get('plateColor'),
                    row.get('vehicleType'),
                    row.get('projectName'),
                    row.get('deptName'),
                    row.get('ownerName'),
                    cls.LEVEL_NAMES.get(row.get('warningLevel'), row.get('warningLevel')),
                    '、'.join(cls.DEVICE_NAMES.get(item, item) for item in row.get('missingDeviceTypes', [])),
                ]
            )
        for column, width in {'A': 16, 'B': 12, 'C': 16, 'D': 22, 'E': 16, 'F': 24, 'G': 12, 'H': 30}.items():
            worksheet.column_dimensions[column].width = width
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
