import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import ColumnElement, select
from sqlalchemy.ext.asyncio import AsyncSession

from module_admin.entity.do.dept_do import SysDept
from module_gat1400.dao.pass_record_dao import TruckPassRecordDao
from module_gat1400.entity.vo.pass_record_vo import PassRecordPageQueryModel
from module_truck.entity.do.truck_do import TruckVehicle

PLATE_COLOR_NAMES = {
    '1': '黑色',
    '2': '白色',
    '5': '蓝色',
    '6': '黄色',
    '9': '绿色',
    '99': '其他',
}


class PassRecordService:
    @classmethod
    async def get_list(
        cls, db: AsyncSession, query: PassRecordPageQueryModel, scope: ColumnElement, is_page: bool
    ) -> Any:
        result = await TruckPassRecordDao.get_list(db, query, scope, is_page)
        rows = result.rows if is_page else result
        vehicle_ids = [row['vehicleId'] for row in rows if row.get('vehicleId')]
        vehicle_depts: dict[int, int] = {}
        if vehicle_ids:
            vehicle_depts = dict(
                (
                    await db.execute(
                        select(TruckVehicle.vehicle_id, TruckVehicle.dept_id).where(
                            TruckVehicle.vehicle_id.in_(vehicle_ids)
                        )
                    )
                ).all()
            )
        dept_names = dict((await db.execute(select(SysDept.dept_id, SysDept.dept_name))).all())
        enriched = []
        for row in rows:
            vehicle_id = row.get('vehicleId')
            dept_id = vehicle_depts.get(vehicle_id) if vehicle_id else None
            enriched.append(
                {
                    'recordId': row['recordId'],
                    'motorVehicleId': row['motorVehicleId'],
                    'plateNo': row.get('plateNo'),
                    'plateColor': row.get('plateColor'),
                    'plateClass': row.get('plateClass'),
                    'vehicleClass': row.get('vehicleClass'),
                    'vehicleColor': row.get('vehicleColor'),
                    'passTime': row.get('passTime'),
                    'nameOfPassedRoad': row.get('nameOfPassedRoad'),
                    'laneNo': row.get('laneNo'),
                    'speed': row.get('speed'),
                    'direction': row.get('direction'),
                    'tollgateId': row.get('tollgateId'),
                    'deviceId': row.get('deviceId'),
                    'storageUrl1': row.get('storageUrl1'),
                    'storageUrl2': row.get('storageUrl2'),
                    'plateReliability': row.get('plateReliability'),
                    'vehicleId': vehicle_id,
                    'vehicleMatched': vehicle_id is not None,
                    'deptName': dept_names.get(dept_id) if dept_id else None,
                }
            )
        if is_page:
            result.rows = enriched
            return result
        return enriched

    @classmethod
    def export_pass_records(cls, rows: list[dict]) -> bytes:
        """过车记录导出：openpyxl 生成，绿色表头，与预警导出风格一致。"""
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '过车记录'
        headers = [
            '车牌号',
            '车牌颜色',
            '车辆类型',
            '车身颜色',
            '过车时间',
            '经过道路',
            '车道',
            '车速(km/h)',
            '行驶方向',
            '卡口编号',
            '设备编码',
            '是否本平台车辆',
            '所属中队',
            '近景照片URL',
        ]
        worksheet.append(headers)
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='2F6B4F')
        for row in rows:
            worksheet.append(
                [
                    row.get('plateNo'),
                    PLATE_COLOR_NAMES.get(str(row.get('plateColor')), row.get('plateColor')),
                    row.get('vehicleClass'),
                    PLATE_COLOR_NAMES.get(str(row.get('vehicleColor')), row.get('vehicleColor')),
                    row.get('passTime').strftime('%Y-%m-%d %H:%M:%S') if row.get('passTime') else '',
                    row.get('nameOfPassedRoad'),
                    row.get('laneNo'),
                    row.get('speed'),
                    row.get('direction'),
                    row.get('tollgateId'),
                    row.get('deviceId'),
                    '是' if row.get('vehicleMatched') else '否',
                    row.get('deptName'),
                    row.get('storageUrl1'),
                ]
            )
        for column, width in {
            'A': 16,
            'B': 12,
            'C': 14,
            'D': 12,
            'E': 20,
            'F': 22,
            'G': 8,
            'H': 12,
            'I': 10,
            'J': 22,
            'K': 22,
            'L': 14,
            'M': 16,
            'N': 45,
        }.items():
            worksheet.column_dimensions[column].width = width
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
